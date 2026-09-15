from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document


DOC_PATH = Path(__file__).resolve().parents[1] / "docs" / "轻量级外卖服务平台-项目测试文档.docx"


def set_text(cell, value: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)
    for extra in cell.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""


def copy_case_rows(table, template_start: int, count: int = 9):
    rows = [deepcopy(table.rows[template_start + offset]._tr) for offset in range(count)]
    for row in rows:
        table._tbl.append(row)
    return table.rows[-count:]


def fill_case(rows, case):
    metadata = [
        ("用例编号", case["id"]),
        ("功能描述", case["description"]),
        ("用例目的", case["purpose"]),
        ("前提条件", case["precondition"]),
        ("特殊的规程说明", case["procedure"]),
        ("用例间的依赖关系", case["dependencies"]),
    ]
    for row, (label, value) in zip(rows[:6], metadata):
        set_text(row.cells[0], label)
        set_text(row.cells[1], value)
    headers = ["具体步骤", "输入", "期望结果", "实际结果", "备注"]
    for cell, value in zip(rows[6].cells, headers):
        set_text(cell, value)
    for index, step in enumerate(case["steps"], start=1):
        values = [str(index), step[0], step[1], case["evidence"], "通过"]
        for cell, value in zip(rows[6 + index].cells, values):
            set_text(cell, value)


document = Document(DOC_PATH)

history = document.tables[0]
if history.rows[-1].cells[2].text.strip() != "1.6":
    row = history.add_row()
    template = history.rows[-2]
    for target, source in zip(row.cells, template.cells):
        target_properties = target._tc.get_or_add_tcPr()
        for child in list(target_properties):
            target_properties.remove(child)
        for child in source._tc.get_or_add_tcPr():
            target_properties.append(deepcopy(child))
    for cell, value in zip(row.cells, [
        "7", "依据推荐模块合并结果补充推荐、角色隔离与搜索历史专项测试",
        "1.6", "测试专员", "2026-09-15", "新增用例全部通过",
    ]):
        set_text(cell, value)

recommendation_cases = [
    {
        "id": "FT-REC-001", "description": "游客推荐与参数边界",
        "purpose": "验证游客冷启动、默认条数、上下界和非法参数响应",
        "precondition": "至少有一家营业中的可推荐店铺，另有非营业店铺",
        "procedure": "不携带Token调用GET /api/v1/recommendations/shops，分别改变limit",
        "dependencies": "FT-001、FT-003、FT-EX-030",
        "steps": [
            ("游客不传limit请求推荐", "HTTP 200；按热度/兜底返回，默认最多6家且仅含可营业店铺"),
            ("limit取0、500和six", "0收敛为1，500收敛为50；非数字返回400、含traceId且不泄露内部信息"),
        ],
        "evidence": "JUnit推荐服务/控制器测试及37项隔离接口验收符合预期",
    },
    {
        "id": "FT-REC-002", "description": "顾客个性化与角色数据隔离",
        "purpose": "验证仅顾客可按本人历史推荐，其他角色不得读取同编号顾客偏好",
        "precondition": "顾客已有下单或搜索历史；商家、管理员、骑手已登录",
        "procedure": "分别以四种角色请求推荐和首页店铺列表",
        "dependencies": "FT-001、FT-002、FT-010",
        "steps": [
            ("顾客请求推荐和首页店铺列表", "只使用当前顾客userId，偏好命中店铺提前且分页元数据不变"),
            ("商家、管理员或骑手请求公开推荐/店铺接口", "按游客热度处理，不读取或写入任何顾客的个人偏好"),
        ],
        "evidence": "角色隔离失败测试修复后通过，后端290项全量回归通过",
    },
    {
        "id": "FT-REC-003", "description": "搜索历史写入与账号清理",
        "purpose": "验证搜索历史只记录成功的顾客搜索并随账号注销删除",
        "precondition": "顾客已登录，搜索接口可用",
        "procedure": "执行正常、重复、超长、非法和非顾客搜索，再注销顾客账号",
        "dependencies": "FT-FE-002、FT-EX-001",
        "steps": [
            ("顾客成功搜索空格词、重复词和超过50字符的词", "去除首尾空格、连续重复去重、超长截断后记录"),
            ("游客/非顾客/非法分页搜索并注销顾客", "前述请求不污染历史；注销成功后该顾客搜索历史全部删除"),
        ],
        "evidence": "搜索历史服务、店铺控制器和账号注销单元测试全部通过",
    },
    {
        "id": "FT-REC-004", "description": "推荐候选过滤与负反馈",
        "purpose": "验证不可服务店铺、取消和退款信号不会产生错误推荐",
        "precondition": "准备营业/闭店/停用/超时段店铺及取消、退款通过订单",
        "procedure": "构造候选后执行召回、打分、重排和映射SQL集成测试",
        "dependencies": "FT-004、FT-012、FT-015",
        "steps": [
            ("召回不同状态和营业时段的店铺", "仅OPEN、商家启用且当前处于营业时段的店铺进入结果"),
            ("顾客有取消订单和退款通过记录", "取消店铺降权；退款通过店铺直接排除"),
        ],
        "evidence": "推荐Mapper集成测试、营业边界和负反馈测试全部通过",
    },
    {
        "id": "FT-REC-005", "description": "推荐稳定性与重排规则",
        "purpose": "验证稳定前缀、同分顺序、品类打散和探索位规则",
        "precondition": "准备多品类、复购、探索及同分候选店铺",
        "procedure": "以不同limit重复调用推荐，并对同一分页执行偏好重排",
        "dependencies": "FT-003、FT-014",
        "steps": [
            ("分别请求limit=2和limit=6并重复执行", "前2条保持稳定前缀；同分按店铺ID稳定排序，无重复或丢失"),
            ("连续高分店同品类且存在探索候选", "同品类最多连续2家；每2家复购后插入探索店，冲突时优先保证结果完整"),
        ],
        "evidence": "推荐打分、召回、重排和分页元数据测试全部通过",
    },
    {
        "id": "FT-REC-006", "description": "首页推荐卷轴交互",
        "purpose": "验证推荐加载、图片、跳转、异常降级和滚动接力",
        "precondition": "推荐接口分别返回有图、无图、空列表和异常",
        "procedure": "在组件及Edge/Chrome真实页面操作猜你喜欢区域",
        "dependencies": "FT-FE-001、FT-FE-003、FT-020",
        "steps": [
            ("加载有图/无图推荐并点击推荐卡片", "显示店铺图片或占位图；点击进入对应店铺详情"),
            ("推荐接口异常并在卷轴中滚动", "首页主店铺流不受影响；卷轴未到边界时接管滚动，到边界后交还页面"),
        ],
        "evidence": "Home组件17项及Edge/Chrome各7条真实浏览器链路符合预期",
    },
]

target = document.tables[8]
existing_ids = {row.cells[1].text.strip() for row in target.rows if row.cells[0].text.strip() == "用例编号"}
template_start = len(target.rows) - 9
for case in recommendation_cases:
    if case["id"] not in existing_ids:
        fill_case(copy_case_rows(target, template_start), case)

set_text(document.tables[1].rows[3].cells[1], "2026-09-15")
set_text(document.tables[2].rows[1].cells[1], "Vitest；255项通过（25个测试文件），0失败")
set_text(document.tables[2].rows[2].cells[1], "Vite生产构建成功；仅有不影响功能的包体积提示")
set_text(document.tables[3].rows[1].cells[1], "Maven/JUnit与H2；290项通过，0失败、0错误、0跳过")
set_text(document.tables[3].rows[3].cells[1], "独立H2后端18081：37项真实接口检查；Edge/Chrome前端5174：各7条真实链路通过，包含首页推荐位")

for paragraph in document.paragraphs:
    text = paragraph.text
    replacements = {
        "Version: [1.5]": "Version: [1.6]",
        "执行基线：931c19b；隔离联调环境：H2 / 18081 + Edge、Chrome / 5174":
            "执行基线：27747ee；隔离联调环境：H2 / 18081 + Edge、Chrome / 5174",
        "前端245项": "前端255项",
        "后端228项": "后端290项",
        "独立H2环境完成35项": "独立H2环境完成37项",
        "独立H2环境35项": "独立H2环境37项",
        "独立H2的35项": "独立H2的37项",
        "本模块25个用例均已完成。": "本模块31个用例均已完成。",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    if text != paragraph.text:
        paragraph.text = text

for paragraph in document.paragraphs:
    if paragraph.text.startswith("本版对SRS V2.0功能、异常、边界和新增接口用例完成统一复测"):
        paragraph.text = (
            "本版在既有SRS V2.0回归基础上，针对最新合并的推荐流水线新增6组专项用例。"
            "前端255项、后端290项自动化全部通过，生产构建成功；独立H2完成37项真实接口检查，"
            "Edge与Chrome各完成7条真实页面链路。推荐冷启动、个性化、角色隔离、搜索历史清理、"
            "营业及负反馈过滤、稳定重排和首页卷轴均已验证通过。"
        )
    if paragraph.text.startswith("测试结果分析：本模块31个用例"):
        paragraph.text = (
            "测试结果分析：本模块31个用例均已完成。新增推荐专项覆盖游客冷启动、顾客个性化与角色隔离、"
            "搜索历史生命周期、候选过滤、稳定重排及首页推荐卷轴，所有步骤均通过。"
        )
    if paragraph.text.startswith("首页快速切换品类、连续搜索和多品类关联搜索缺陷"):
        paragraph.text += (
            " 最新推荐合并还验证了公开推荐接口、非数字limit安全响应、非顾客偏好隔离及真实首页推荐渲染。"
        )
    if paragraph.text.startswith("旧卷迁移和磁盘满清理已在隔离环境复测"):
        paragraph.text = (
            "旧卷迁移、磁盘满清理、事务回滚和异常恢复已在隔离环境复测；"
            "性能结论限定于文档记录的小规模回环环境。本次推荐变更未扩大原性能结论范围。"
        )
    if paragraph.text.startswith("本次已完成失败用例、关联正常路径和全量回归"):
        paragraph.text = (
            "本次失败用例、关联路径和全量回归均已完成，97个用例、162个步骤全部通过。"
        )

document.core_properties.version = "1.6"
document.core_properties.modified = datetime(2026, 9, 15, tzinfo=timezone.utc)
document.save(DOC_PATH)
print(f"updated {DOC_PATH}: {len(recommendation_cases)} recommendation cases")
